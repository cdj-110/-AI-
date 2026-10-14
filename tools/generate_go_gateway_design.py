from __future__ import annotations

import os
from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "交付文档"
ASSET_DIR = OUT_DIR / "go网关系统设计说明书_assets"
OUT_FILE = OUT_DIR / "微控Go工业网关系统设计说明书_V1.0.docx"

# Design preset: standard_business_brief.
# Named override: Microsoft YaHei is used in place of Calibri for CJK legibility.
FONT = "Microsoft YaHei"
FONT_EN = "Calibri"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "667085"
LIGHT = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
VERY_LIGHT_BLUE = "F5F9FF"
GREEN = "12A150"
ORANGE = "E8871E"
RED = "C0392B"
WHITE = "FFFFFF"
BLACK = "111827"
BORDER = "D8E0EA"


def set_run_font(run, name=FONT, size=11, color=BLACK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_inches, indent_dxa=120):
    total_dxa = int(round(sum(widths_inches) * 1440))
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(round(width * 1440))))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_inches):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(round(width * 1440))))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT


def add_table(doc, headers, rows, widths, header_fill=LIGHT, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_shading(header[i], header_fill)
        p = header[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(str(text)), size=font_size, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if i == 0 and len(headers) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(str(value)), size=font_size, color=BLACK)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run(text), size=9, color=MUTED)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text))
    return p


def add_callout(doc, title, text, fill=VERY_LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(title + "  "), size=10.2, color=accent, bold=True)
    set_run_font(p.add_run(text), size=10.2, color=BLACK)
    set_table_geometry(table, [6.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run("微控 Go 工业网关 | 系统设计说明书"), size=9, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(p)


def font(size=30, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def wrap_text(draw, text, fnt, max_width):
    lines, current = [], ""
    for ch in text:
        candidate = current + ch
        if current and draw.textbbox((0, 0), candidate, font=fnt)[2] > max_width:
            lines.append(current)
            current = ch
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def rounded_box(draw, xy, title, lines=(), fill="#F7FAFC", outline="#B8C8DB", title_fill=None, title_color="#143D63"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    if title_fill:
        draw.rounded_rectangle((x1, y1, x2, min(y1 + 54, y2)), radius=18, fill=title_fill)
        draw.rectangle((x1, y1 + 30, x2, min(y1 + 54, y2)), fill=title_fill)
    tf = font(26, True)
    bf = font(20)
    tw = draw.textbbox((0, 0), title, font=tf)[2]
    draw.text(((x1 + x2 - tw) / 2, y1 + 12), title, font=tf, fill=title_color)
    y = y1 + 70
    for line in lines:
        for part in wrap_text(draw, line, bf, x2 - x1 - 34):
            draw.text((x1 + 18, y), part, font=bf, fill="#364152")
            y += 29
        y += 4


def arrow(draw, start, end, color="#4777A8", width=4):
    draw.line((start, end), fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 15
    for delta in (2.55, -2.55):
        p = (end[0] + length * math.cos(angle + delta), end[1] + length * math.sin(angle + delta))
        draw.line((end, p), fill=color, width=width)


def canvas(title, width=1600, height=900):
    im = Image.new("RGB", (width, height), "#FFFFFF")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((12, 12, width - 12, height - 12), radius=24, fill="#FFFFFF", outline="#D8E0EA", width=3)
    d.text((44, 28), title, font=font(34, True), fill="#0B2545")
    d.line((44, 82, width - 44, 82), fill="#D8E0EA", width=2)
    return im, d


def save_architecture():
    im, d = canvas("系统总体结构")
    boxes = {
        "field": (55, 150, 300, 760),
        "collect": (355, 145, 640, 415),
        "runtime": (700, 145, 985, 415),
        "state": (700, 500, 985, 770),
        "services": (1045, 145, 1340, 770),
        "external": (1395, 150, 1545, 760),
    }
    rounded_box(d, boxes["field"], "现场层", ["Modbus RTU/TCP", "Siemens S7", "IEC 104", "OPC UA", "IEC 61850 MMS", "IEC 61850 GOOSE"], "#F5F9FF", "#9FC3E8", "#DDEEFF")
    rounded_box(d, boxes["collect"], "协议采集层", ["Collector 工厂", "连接复用与批量读取", "通道分 lane 调度", "报文监控"], "#F8FBFF", "#8FB9E2", "#DDEEFF")
    rounded_box(d, boxes["runtime"], "运行编排层", ["runtime.Manager", "配置热加载", "采集退避/写优先", "点位映射与状态更新"], "#F6FAF7", "#8DC7A5", "#DFF3E7", "#17653A")
    rounded_box(d, boxes["state"], "实时数据层", ["state.Store", "内存点位索引", "revision + diff", "WebSocket 合并推送"], "#F6FAF7", "#8DC7A5", "#DFF3E7", "#17653A")
    rounded_box(d, boxes["services"], "业务服务层", ["MQTT / 激活连接", "边缘计算 CEL", "协议转发服务", "离线缓存 / 历史", "Web API / AI 助手", "审计 / 网络 / 维护"], "#FFF9F2", "#E9BE87", "#FFE9CE", "#8B4A00")
    rounded_box(d, boxes["external"], "外部", ["Vue 浏览器", "微控云", "第三方 MQTT", "下级采集端"], "#FBF7FF", "#C4A8DE", "#EFE2F8", "#6F3A8C")
    arrow(d, (300, 330), (355, 280))
    arrow(d, (640, 280), (700, 280))
    arrow(d, (840, 415), (840, 500))
    arrow(d, (985, 280), (1045, 280))
    arrow(d, (1045, 600), (985, 640))
    arrow(d, (1340, 330), (1395, 330))
    arrow(d, (1395, 560), (1340, 560))
    path = ASSET_DIR / "01_system_architecture.png"
    im.save(path, quality=95)
    return path


def save_deployment():
    im, d = canvas("EG100 部署拓扑")
    rounded_box(d, (520, 170, 1080, 730), "EG100 边缘网关", ["Linux / ARMv7", "Go 单进程 + Supervisor", "内嵌 Vue 静态资源", "板载配置与运行目录", "可选 U 盘/TF 历史与缓存"], "#F6FAF7", "#76BA94", "#DFF3E7", "#17653A")
    rounded_box(d, (55, 160, 420, 390), "工业现场", ["RS485-1 / RS485-2", "eth0 / eth1", "GOOSE 二层以太网"], "#F5F9FF", "#8FB9E2", "#DDEEFF")
    rounded_box(d, (55, 500, 420, 730), "运维终端", ["PC / 手机浏览器", "HTTPS / WebSocket", "配置、监控、诊断"], "#FBF7FF", "#C4A8DE", "#EFE2F8")
    rounded_box(d, (1180, 150, 1540, 390), "云平台", ["MQTT Broker", "微控云激活", "远程配置 / 遥测"], "#FFF9F2", "#E9BE87", "#FFE9CE", "#8B4A00")
    rounded_box(d, (1180, 500, 1540, 730), "下游系统", ["Modbus TCP Client", "IEC104 主站", "IEC61850 Client", "GOOSE Subscriber"], "#FFF7F7", "#E5AAAA", "#FCE3E3", "#8C2E2E")
    arrow(d, (420, 275), (520, 300))
    arrow(d, (420, 610), (520, 600))
    arrow(d, (1080, 300), (1180, 275))
    arrow(d, (1080, 610), (1180, 610))
    path = ASSET_DIR / "02_deployment.png"
    im.save(path, quality=95)
    return path


def class_box(draw, x, y, w, name, attrs, methods, color="#DDEEFF"):
    h = 74 + max(1, len(attrs)) * 27 + max(1, len(methods)) * 27
    draw.rounded_rectangle((x, y, x + w, y + h), radius=14, fill="#FFFFFF", outline="#7EA8D2", width=3)
    draw.rounded_rectangle((x, y, x + w, y + 54), radius=14, fill=color)
    draw.rectangle((x, y + 32, x + w, y + 54), fill=color)
    tf, bf = font(24, True), font(18)
    draw.text((x + 16, y + 12), name, font=tf, fill="#143D63")
    yy = y + 62
    for item in attrs:
        draw.text((x + 14, yy), "- " + item, font=bf, fill="#475467")
        yy += 27
    draw.line((x, yy + 2, x + w, yy + 2), fill="#D8E0EA", width=2)
    yy += 10
    for item in methods:
        draw.text((x + 14, yy), "+ " + item, font=bf, fill="#1F4D78")
        yy += 27
    return (x, y, x + w, y + h)


def save_class_core():
    im, d = canvas("核心运行类图")
    app = class_box(d, 60, 145, 340, "app.App", ["cfg Config", "state *Store", "runtime *Manager", "clouds []*Client"], ["Run(ctx)", "ApplyConfig(cfg)", "publishToCloud(batch)"], "#DDEEFF")
    runtime = class_box(d, 470, 145, 340, "runtime.Manager", ["cfg Config", "lanes map", "retryAfter map"], ["CollectLane()", "WritePoint()", "UpdateConfig()"], "#DFF3E7")
    state = class_box(d, 880, 145, 320, "state.Store", ["points map", "revisions", "subscribers"], ["SetPointValue()", "Snapshot()", "Subscribe()"], "#DFF3E7")
    web = class_box(d, 1270, 145, 270, "web.Server", ["runtime", "data", "audit", "ai"], ["Run(ctx)", "routes()", "webSocket()"], "#EFE2F8")
    coll = class_box(d, 265, 540, 310, "collector.Collector", ["protocol adapter"], ["ReadPoint()", "WritePoint()"], "#FFE9CE")
    edge = class_box(d, 650, 540, 300, "edgecompute.Manager", ["program *Program", "output chan"], ["Start(ctx)", "evaluate()", "UpdateConfig()"], "#FFE9CE")
    fwd = class_box(d, 1025, 540, 310, "forward.Manager", ["cfg Config", "store *Store"], ["Update(ctx,cfg)", "serveModbus()", "serveIEC104()"], "#FFE9CE")
    arrow(d, (400, 270), (470, 270)); arrow(d, (810, 270), (880, 270)); arrow(d, (1200, 270), (1270, 270))
    arrow(d, (560, 415), (470, 540)); arrow(d, (740, 420), (800, 540)); arrow(d, (1010, 420), (1170, 540))
    path = ASSET_DIR / "03_core_class.png"
    im.save(path, quality=95)
    return path


def save_class_domain():
    im, d = canvas("配置领域类图")
    root = class_box(d, 625, 120, 350, "config.Config", ["resources []Resource", "channels []Channel", "devices []Device", "points []Point"], ["ApplyDefaults()", "Validate()", "Save()"], "#DDEEFF")
    res = class_box(d, 50, 460, 270, "ResourceConfig", ["resourceKey", "type", "serial/network"], ["物理资源"], "#DFF3E7")
    channel = class_box(d, 370, 460, 270, "ChannelConfig", ["channelKey", "resourceKey", "role/protocol"], ["采集或转发通道"], "#DFF3E7")
    dev = class_box(d, 690, 460, 270, "DeviceConfig", ["deviceKey", "channelKey", "address/protocol"], ["IsEnabled()"], "#DFF3E7")
    point = class_box(d, 1010, 460, 270, "PointConfig", ["deviceKey", "metric", "protocol fields"], ["CollectInterval()"], "#DFF3E7")
    forward = class_box(d, 1330, 460, 220, "ForwardDevice", ["deviceKey", "points []Mapping"], ["IsEnabled()"], "#FFE9CE")
    for x in (185, 505, 825, 1145, 1440):
        arrow(d, (800, root[3]), (x, 460))
    arrow(d, (320, 590), (370, 590)); arrow(d, (640, 590), (690, 590)); arrow(d, (960, 590), (1010, 590))
    path = ASSET_DIR / "04_domain_class.png"
    im.save(path, quality=95)
    return path


def save_er():
    im, d = canvas("配置数据概念 ER 图")
    entities = {
        "gateway": (640, 115, 960, 260, "GatewayConfig", ["PK gatewayKey", "schemaVersion"]),
        "resource": (60, 350, 330, 535, "Resource", ["PK resourceKey", "type", "enabled"]),
        "channel": (390, 350, 660, 535, "Channel", ["PK channelKey", "FK resourceKey", "role/protocol"]),
        "device": (720, 350, 990, 535, "Device", ["PK deviceKey", "FK channelKey", "address"]),
        "point": (1050, 350, 1320, 535, "Point", ["PK deviceKey+metric", "FK channelKey", "dataType"]),
        "mqtt": (1380, 350, 1545, 535, "MQTTChannel", ["name", "broker", "clientId"]),
        "forward": (120, 665, 430, 840, "ForwardDevice", ["PK deviceKey", "FK channelKey"]),
        "mapping": (505, 665, 815, 840, "ForwardPoint", ["PK deviceKey+metric", "FK source point"]),
        "edge": (890, 665, 1180, 840, "EdgeComputedPoint", ["PK groupKey+metric", "expression"]),
        "user": (1260, 665, 1545, 840, "User / Role", ["PK username / roleKey", "FK user.roleKey"]),
    }
    for _, (x1, y1, x2, y2, title, attrs) in entities.items():
        rounded_box(d, (x1, y1, x2, y2), title, attrs, "#FFFFFF", "#8EABC7", "#E8EEF5")
    for x in (195, 525, 855, 1185, 1460): arrow(d, (800, 260), (x, 350))
    arrow(d, (330, 445), (390, 445)); arrow(d, (660, 445), (720, 445)); arrow(d, (990, 445), (1050, 445))
    arrow(d, (275, 665), (525, 535)); arrow(d, (660, 665), (1185, 535)); arrow(d, (1035, 665), (1185, 535))
    d.text((338, 410), "1 : N", font=font(18, True), fill="#4777A8")
    d.text((668, 410), "1 : N", font=font(18, True), fill="#4777A8")
    d.text((998, 410), "1 : N", font=font(18, True), fill="#4777A8")
    path = ASSET_DIR / "05_er.png"
    im.save(path, quality=95)
    return path


def save_sequence():
    im, d = canvas("采集、计算与上报时序")
    actors = ["调度器", "runtime", "Collector", "state.Store", "边缘计算", "MQTT"]
    xs = [135, 390, 650, 920, 1180, 1450]
    for x, actor in zip(xs, actors):
        d.rounded_rectangle((x - 92, 115, x + 92, 170), radius=12, fill="#E8EEF5", outline="#8EAED0", width=2)
        tw = d.textbbox((0, 0), actor, font=font(22, True))[2]
        d.text((x - tw / 2, 129), actor, font=font(22, True), fill="#143D63")
        d.line((x, 170, x, 820), fill="#C7D2DF", width=2)
    events = [
        (240, 0, 1, "到期 lane"), (310, 1, 2, "批量/单点读取"), (380, 2, 1, "PointValue / error"),
        (450, 1, 3, "SetPointValue"), (520, 3, 4, "revision 通知"), (590, 4, 3, "写入派生点"),
        (660, 1, 5, "差异遥测 / 全量模板"), (730, 5, 1, "PUBACK / 失败"), (790, 1, 3, "MarkPublish / 缓存"),
    ]
    for y, a, b, label in events:
        arrow(d, (xs[a], y), (xs[b], y), "#4777A8" if b >= a else "#78909C", 3)
        mid = (xs[a] + xs[b]) / 2
        tw = d.textbbox((0, 0), label, font=font(18))[2]
        d.text((mid - tw / 2, y - 25), label, font=font(18), fill="#364152")
    path = ASSET_DIR / "06_sequence.png"
    im.save(path, quality=95)
    return path


def wireframe_shell(draw, active="运行状态"):
    draw.rectangle((45, 100, 275, 850), fill="#F7F9FC", outline="#D8E0EA")
    draw.text((72, 125), "微控网关", font=font(28, True), fill="#0B2545")
    items = ["运行状态", "智能网关", "边缘计算", "AI 助手", "WiFi", "云平台连接", "系统维护"]
    y = 205
    for item in items:
        fill = "#E7F1FF" if item == active else "#F7F9FC"
        draw.rounded_rectangle((58, y, 258, y + 54), radius=10, fill=fill)
        draw.text((82, y + 13), item, font=font(20, item == active), fill="#1769CC" if item == active else "#3F4B5A")
        y += 68
    draw.rectangle((275, 100, 1555, 175), fill="#FFFFFF", outline="#D8E0EA")
    draw.text((310, 121), active, font=font(30, True), fill="#111827")
    draw.rounded_rectangle((1360, 118, 1515, 158), radius=12, fill="#E8EEF5")
    draw.text((1392, 124), "配置  运行", font=font(18, True), fill="#143D63")


def save_prototype_status():
    im, d = canvas("原型：运行状态（桌面端）")
    wireframe_shell(d, "运行状态")
    x, y = 310, 205
    for label, value, color in (("点位总数", "1008", "#1769CC"), ("正常点位", "1002", "#12A150"), ("异常点位", "6", "#C0392B"), ("内存占用", "27.3 MB", "#111827")):
        d.rounded_rectangle((x, y, x + 275, y + 120), radius=16, fill="#FFFFFF", outline="#D8E0EA", width=2)
        d.text((x + 20, y + 18), label, font=font(19), fill="#667085")
        d.text((x + 20, y + 52), value, font=font(32, True), fill=color)
        x += 300
    d.rounded_rectangle((310, 355, 1515, 535), radius=18, fill="#FFFFFF", outline="#D8E0EA", width=2)
    d.text((335, 378), "实时指标", font=font(25, True), fill="#111827")
    for i, (name, val) in enumerate((("温度", "26.4 ℃"), ("压力", "0.82 MPa"), ("开关状态", "合位"), ("功率", "18.6 kW"))):
        xx = 335 + i * 285
        d.rounded_rectangle((xx, 425, xx + 250, 505), radius=12, fill="#F5F9FF", outline="#C9D9EC")
        d.text((xx + 14, 438), name, font=font(18), fill="#667085")
        d.text((xx + 14, 467), val, font=font(24, True), fill="#1769CC")
    d.rounded_rectangle((310, 560, 1515, 825), radius=18, fill="#FFFFFF", outline="#D8E0EA", width=2)
    d.text((335, 582), "点位状态 / WebSocket 实时更新", font=font(24, True), fill="#111827")
    headers = ["子设备", "点位", "协议", "当前值", "更新时间", "状态"]
    for i, h in enumerate(headers): d.text((335 + i * 190, 635), h, font=font(18, True), fill="#667085")
    for r in range(3):
        yy = 680 + r * 42
        vals = ["device-001", f"mbtcp-P0{r+1}", "modbus-tcp", str(120 + r * 7), "14:05:2" + str(r), "正常"]
        for i, v in enumerate(vals): d.text((335 + i * 190, yy), v, font=font(17), fill="#12A150" if i == 5 else "#344054")
    path = ASSET_DIR / "07_prototype_status.png"
    im.save(path, quality=95)
    return path


def save_prototype_config():
    im, d = canvas("原型：智能网关设备与点位配置（桌面端）")
    wireframe_shell(d, "智能网关")
    d.rounded_rectangle((310, 205, 1515, 335), radius=16, fill="#FFFFFF", outline="#D8E0EA", width=2)
    d.text((335, 225), "设备1", font=font(28, True), fill="#111827")
    d.text((335, 270), "所属通道：网口1 / Modbus TCP 采集", font=font(19), fill="#667085")
    d.rounded_rectangle((1250, 238, 1475, 292), radius=12, fill="#EEF6FF", outline="#9FC3E8")
    d.text((1280, 251), "设备配置 | 点位配置", font=font(18, True), fill="#1769CC")
    d.rounded_rectangle((310, 355, 1515, 825), radius=16, fill="#FFFFFF", outline="#D8E0EA", width=2)
    d.text((335, 378), "点位管理", font=font(24, True), fill="#111827")
    for i, t in enumerate(("线圈 01", "离散输入 02", "保持寄存器 03", "输入寄存器 04")):
        xx = 335 + i * 220
        fill = "#E7F1FF" if i == 0 else "#FFFFFF"
        d.rounded_rectangle((xx, 420, xx + 200, 470), radius=10, fill=fill, outline="#AFC8E2")
        d.text((xx + 18, 432), t, font=font(18, i == 0), fill="#1769CC" if i == 0 else "#475467")
    d.rounded_rectangle((335, 495, 1490, 550), radius=10, fill="#F7F9FC", outline="#D8E0EA")
    d.text((355, 510), "+ 新增点位    复制选中    批量修改    删除选中", font=font(18), fill="#344054")
    d.text((1260, 510), "实时报文监控", font=font(18, True), fill="#1769CC")
    headers = ["#", "名称", "标识符", "实时值", "寄存器", "数据类型", "倍率", "操作"]
    widths = [55, 285, 215, 135, 125, 150, 100, 110]
    xx = 340
    for h, w in zip(headers, widths):
        d.text((xx, 585), h, font=font(17, True), fill="#667085"); xx += w
    for r in range(4):
        yy = 630 + r * 43
        vals = [str(r+1), f"eth0_Modbus采集_设备1@dy{r+1}", f"mbtcp-P0{r+1}", str(110+r*12), str(r), "uint16", "1", "下发"]
        xx = 340
        for v, w in zip(vals, widths):
            d.text((xx, yy), v, font=font(16), fill="#1769CC" if v == "下发" else "#344054"); xx += w
        d.line((335, yy+31, 1490, yy+31), fill="#E6EBF0", width=1)
    path = ASSET_DIR / "08_prototype_config.png"
    im.save(path, quality=95)
    return path


def save_prototype_ai():
    im, d = canvas("原型：AI 助手与 PDF 配置生成（桌面端）")
    wireframe_shell(d, "AI 助手")
    d.text((315, 215), "智能对话     PDF 生成配置", font=font(23, True), fill="#1769CC")
    d.rounded_rectangle((310, 270, 1515, 715), radius=18, fill="#FFFFFF", outline="#D8E0EA", width=2)
    d.rounded_rectangle((355, 310, 1030, 390), radius=18, fill="#F3F7FC")
    d.text((380, 330), "AI：我可以读取脱敏配置并协助新增通道、设备和点位。", font=font(20), fill="#344054")
    d.rounded_rectangle((875, 420, 1465, 490), radius=18, fill="#E8F3FF")
    d.text((900, 440), "用户：为网口1创建一个 IEC104 采集设备", font=font(20), fill="#1769CC")
    d.rounded_rectangle((355, 520, 1160, 650), radius=16, fill="#F6FAF7", outline="#9BD2B1")
    d.text((380, 540), "需要确认", font=font(20, True), fill="#128A48")
    d.text((380, 578), "网口资源：网口1 (eth0)     模式：新增设备", font=font(18), fill="#344054")
    d.rounded_rectangle((960, 590, 1125, 635), radius=10, fill="#1677FF")
    d.text((1000, 600), "确认执行", font=font(18, True), fill="#FFFFFF")
    d.rounded_rectangle((310, 740, 1350, 815), radius=12, fill="#FFFFFF", outline="#B9C7D8")
    d.text((335, 760), "输入网关相关问题；Enter 发送，Shift+Enter 换行", font=font(19), fill="#98A2B3")
    d.rounded_rectangle((1370, 740, 1515, 815), radius=12, fill="#1677FF")
    d.text((1414, 760), "发送", font=font(20, True), fill="#FFFFFF")
    path = ASSET_DIR / "09_prototype_ai.png"
    im.save(path, quality=95)
    return path


def save_prototype_mobile():
    im, d = canvas("原型：移动端查看优先布局", width=1050, height=1500)
    d.rounded_rectangle((210, 100, 840, 1430), radius=45, fill="#F5F7FA", outline="#AAB8C8", width=5)
    d.rectangle((235, 145, 815, 250), fill="#FFFFFF")
    d.text((280, 178), "☰", font=font(34, True), fill="#344054")
    d.text((365, 172), "智能网关", font=font(32, True), fill="#111827")
    d.text((750, 172), "⚙", font=font(30), fill="#667085")
    d.rounded_rectangle((255, 280, 795, 390), radius=18, fill="#FFFFFF", outline="#D8E0EA")
    d.text((280, 300), "设备1  ·  在线", font=font(27, True), fill="#111827")
    d.text((280, 344), "Modbus TCP  |  13 个点位  |  1 秒刷新", font=font(19), fill="#667085")
    d.rounded_rectangle((255, 415, 795, 475), radius=15, fill="#E8F3FF")
    d.text((290, 430), "设备参数（已折叠）                       ˅", font=font(20, True), fill="#1769CC")
    d.text((260, 515), "实时数据", font=font(27, True), fill="#111827")
    d.rounded_rectangle((255, 560, 795, 625), radius=15, fill="#FFFFFF", outline="#D8E0EA")
    d.text((285, 578), "搜索点位名称或标识符", font=font(19), fill="#98A2B3")
    cards = [("温度", "26.4", "℃", "#1769CC"), ("运行", "开启", "", "#12A150"), ("压力", "0.82", "MPa", "#1769CC"), ("告警", "正常", "", "#12A150")]
    for i, (name, value, unit, color) in enumerate(cards):
        col, row = i % 2, i // 2
        x, y = 255 + col * 275, 660 + row * 230
        d.rounded_rectangle((x, y, x + 260, y + 205), radius=18, fill="#FFFFFF", outline="#D8E0EA")
        d.text((x + 20, y + 20), name, font=font(21, True), fill="#344054")
        d.text((x + 20, y + 82), value, font=font(40, True), fill=color)
        d.text((x + 160, y + 100), unit, font=font(19), fill="#667085")
        d.text((x + 20, y + 160), "刚刚更新", font=font(17), fill="#98A2B3")
    d.rectangle((235, 1300, 815, 1400), fill="#FFFFFF", outline="#D8E0EA")
    for x, label in zip((300, 430, 560, 690), ("状态", "网关", "边缘", "更多")):
        d.text((x, 1337), label, font=font(18, label == "网关"), fill="#1769CC" if label == "网关" else "#667085")
    path = ASSET_DIR / "10_prototype_mobile.png"
    im.save(path, quality=95)
    return path


def add_image(doc, path, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    return p


def build_document(diagrams):
    doc = Document()
    configure_document(doc)

    # Cover / memo masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(45)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("WEIKONG EDGE GATEWAY"), size=11, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("微控 Go 工业网关"), size=29, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    set_run_font(p.add_run("系统设计说明书"), size=23, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    set_run_font(p.add_run("面向 EG100 的采集、转发、边缘计算、云连接与本地运维一体化设计"), size=12.5, color=MUTED)
    add_table(doc, ["文档属性", "内容"], [
        ("文档版本", "V1.0"),
        ("实现基线", "weikong-iot-platform / apps/gateway-go（2026-08-20）"),
        ("目标硬件", "EG100（Linux / ARMv7，兼容 Windows 本地测试构建）"),
        ("适用对象", "研发、测试、实施、运维、技术评审与项目交付人员"),
        ("保密等级", "内部技术资料"),
    ], [1.55, 4.95], font_size=10)
    add_callout(doc, "文档定位", "本说明书以当前 Go 网关代码为事实基线，描述已实现架构，同时对配置数据建立概念 ER 模型。网关未引入关系型数据库，ER 图用于表达实体关系，不代表物理部署了 MySQL。")
    page_break(doc)

    heading(doc, "目录", 1)
    toc = [
        "1. 文档概述", "2. 系统总体设计", "3. 系统总体结构", "4. 核心模块设计",
        "5. 系统类图与说明", "6. 数据库与持久化设计", "7. 核心流程与接口设计",
        "8. 安全、可靠性与性能设计", "9. 原型设计", "10. 部署与运行设计", "11. 验证与验收建议", "附录 A. 术语与代码映射",
    ]
    for item in toc:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(item), size=11, color=DARK_BLUE, bold=item[0].isdigit())

    heading(doc, "1. 文档概述", 1)
    heading(doc, "1.1 编写目的", 2)
    add_body(doc, "本文档用于统一描述微控 Go 工业网关的系统边界、总体结构、核心对象、数据持久化、接口、原型与部署方式，为研发迭代、代码评审、测试设计、现场实施和后续维护提供一致依据。")
    heading(doc, "1.2 系统范围", 2)
    for text in [
        "现场采集：Modbus TCP/RTU、Siemens S7、IEC 104、OPC UA、IEC 61850 MMS 和 IEC 61850 GOOSE。",
        "协议转发：Modbus TCP 从站、IEC 104 服务端、IEC 61850 MMS 服务端和 GOOSE 发布。",
        "本地能力：实时状态、WebSocket 差异推送、点位下发、边缘计算、报文监控、网络配置、用户权限、审计和维护。",
        "云连接：微控云激活连接和多个独立 MQTT 通道，支持扁平键值、按设备分组与自定义 JSON。",
        "智能辅助：DeepSeek 对话操作和百炼 PDF OCR 生成配置草稿，所有配置变更需确认并校验。",
    ]:
        add_bullet(doc, text)
    heading(doc, "1.3 设计约束", 2)
    add_table(doc, ["约束", "设计响应"], [
        ("嵌入式资源有限", "Go 单进程、受控并发、差异推送、队列合并、超大点表窗口化加载"),
        ("现场协议差异大", "Collector 工厂 + 协议专用适配器 + 统一 PointConfig"),
        ("不能因页面或云端阻塞采集", "采集、WebSocket、手工 MQTT、激活 MQTT 分离；慢消费者不反压采集"),
        ("掉电和配置错误可恢复", "原子保存、previous 备份、热加载失败回滚、Supervisor/Watchdog"),
        ("默认无外接存储", "历史与离线缓存仅在可移动介质可用时启用"),
    ], [1.65, 4.85])

    page_break(doc)
    heading(doc, "2. 系统总体设计", 1)
    heading(doc, "2.1 设计目标", 2)
    add_table(doc, ["目标", "说明", "关键机制"], [
        ("实时性", "毫秒到秒级采集与秒级上报", "按设备/通道 lane 调度、批量寄存器读取、独立 MQTT 节拍"),
        ("稳定性", "单协议故障不拖垮整体", "连接退避、超时、隔离 goroutine、健康检查"),
        ("可维护性", "统一配置模型和页面操作", "Config 校验、REST/WebSocket API、工程导入导出"),
        ("可扩展性", "新增采集/转发协议代价可控", "Collector/PointWriter 接口与 Forward Manager"),
        ("大点表", "支持 3 万至 30 万级配置管理", "compact API、虚拟表格、可见窗口实时值、增量变更"),
    ], [1.05, 2.15, 3.30])
    heading(doc, "2.2 分层原则", 2)
    add_number(doc, "配置层只描述业务与协议参数，不直接持有网络连接。")
    add_number(doc, "运行层将点位编排为采集 lane，并负责退避、超时和写优先级。")
    add_number(doc, "状态层只保存最新状态和修订号，不让浏览器读取路径影响采集。")
    add_number(doc, "业务服务层以状态层为数据源完成边缘计算、转发、上报和历史记录。")
    add_number(doc, "Web 层通过权限、ETag、审计和回滚控制配置变更。")

    heading(doc, "3. 系统总体结构", 1)
    add_image(doc, diagrams["architecture"])
    add_caption(doc, "图 3-1  Go 网关总体分层架构")
    add_body(doc, "现场协议数据首先进入协议采集适配器，runtime.Manager 按采集 lane 调度并将结果写入 state.Store。状态变化同时驱动 WebSocket 差异推送、边缘计算、协议转发和 MQTT 上报。历史与离线缓存位于持久化边界，外部不可用时不会反向阻塞现场采集。")
    add_image(doc, diagrams["deployment"])
    add_caption(doc, "图 3-2  EG100 部署与外部系统拓扑")

    page_break(doc)
    heading(doc, "4. 核心模块设计", 1)
    add_table(doc, ["模块", "职责", "关键实现"], [
        ("config", "配置加载、默认值、校验、脱敏、原子保存", "Config、Save、PreserveMaskedSecrets"),
        ("collector", "统一点位读取/写入和协议连接复用", "Collector、PointWriter、RegisterRangeReader"),
        ("runtime", "采集 lane、到期调度、退避、批量读取", "Manager、CollectionLane"),
        ("state", "最新点值、健康统计、修订号与订阅", "Store、PointStatus、Snapshot"),
        ("datamanager", "状态差异、通信差异和日志消息", "Message、PointDiff、DeviceDiff"),
        ("cloud", "激活与手工 MQTT 连接、发布和订阅", "Client、多通道独立状态"),
        ("edgecompute", "CEL 表达式编译、依赖跟踪、派生点输出", "Program、Manager"),
        ("forward", "把采集值映射为下游协议服务", "Modbus/IEC104/MMS/GOOSE"),
        ("web", "Vue 静态资源、REST、WebSocket、认证审计", "Server、Auth、Audit、AI Manager"),
        ("supervisor", "软件看门狗、硬件喂狗和重启限制", "Supervisor、/dev/watchdog"),
    ], [1.15, 2.45, 2.90], font_size=8.9)

    heading(doc, "4.1 采集调度", 2)
    add_body(doc, "点位按 DeviceKey 与 ChannelKey 组合为采集 lane，每个 lane 采用最短配置周期。Modbus 点位按连接和功能码串行，再将连续地址合并为批量请求；跨设备/协议任务通过受控并发执行，避免串口并发冲突和无上限 goroutine。")
    add_callout(doc, "大点表保护", "当单 lane 超过 50,000 点时，运行时按处理能力提高有效周期，防止配置的极短周期造成持续积压。前端大配置模式只加载必要结构与可见点位实时值。")
    heading(doc, "4.2 实时状态与 WebSocket", 2)
    add_body(doc, "state.Store 使用 RWMutex 保护内存状态，以 deviceKey::metric 作为主索引，同时维护稳定显示顺序。变化只写入容量为 1 的合并通知；datamanager 根据 revision 生成 points.diff、gateway.status、communications.diff 和 logs.diff，慢浏览器只接收最新状态，不会阻塞采集线程。")
    heading(doc, "4.3 MQTT 上报", 2)
    add_body(doc, "激活连接与手工 MQTT 通道分别管理。采集产生的遥测按设备和指标合并；手工通道由独立 100 ms 调度器检查各自发布周期，避免某一 Broker 的 QoS 等待拉长其他通道。自定义 JSON 模式支持模板渲染，订阅主题用于下行控制。")
    heading(doc, "4.4 配置热加载", 2)
    add_body(doc, "保存前完成 schema、协议字段、标识符唯一性、边缘表达式和监听地址校验；通过 ETag/If-Match 防止并发覆盖。配置先原子写盘，再执行运行时热加载；若应用失败，恢复 previous 配置并重新应用旧配置。")

    page_break(doc)
    heading(doc, "5. 系统类图与说明", 1)
    add_image(doc, diagrams["class_core"])
    add_caption(doc, "图 5-1  核心运行类图")
    add_table(doc, ["类/接口", "协作说明"], [
        ("app.App", "组合运行时、状态、云连接、缓存、历史、边缘计算和转发，是应用生命周期入口。"),
        ("runtime.Manager", "把 Config 转为 CollectionLane，调用 Collector 采集并更新 State。"),
        ("state.Store", "向 WebSocket、转发和边缘计算提供一致的最新点值视图。"),
        ("web.Server", "提供页面、REST、WebSocket、权限、审计、网络和 AI 接口。"),
        ("collector.Collector", "协议策略接口；每个协议适配器返回统一 PointValue。"),
        ("edgecompute.Manager", "订阅状态变化，执行 CEL 程序并把派生点写回 State。"),
        ("forward.Manager", "读取 State 最新值并对外提供下游服务器/发布者。"),
    ], [1.75, 4.75])
    add_image(doc, diagrams["class_domain"])
    add_caption(doc, "图 5-2  配置领域类图")
    add_body(doc, "Config 是聚合根。Resource 描述串口/网口物理资源，Channel 描述资源上的采集或转发逻辑，Device 描述协议会话，Point 描述具体数据项。ForwardDevice 与 ForwardPoint 引用采集源，不重复采集。")

    page_break(doc)
    heading(doc, "6. 数据库与持久化设计", 1)
    add_callout(doc, "物理实现说明", "当前 Go 网关没有关系型数据库。配置采用 JSON 文档，实时数据保存在内存，离线缓存、历史和审计采用 JSONL 文件。以下 ER 图是配置领域的概念数据模型，用于约束主键、外键和唯一性。", "FFF9F2", ORANGE)
    add_image(doc, diagrams["er"])
    add_caption(doc, "图 6-1  配置数据概念 ER 图")
    heading(doc, "6.1 实体与约束", 2)
    add_table(doc, ["实体", "逻辑主键", "主要关系与约束"], [
        ("GatewayConfig", "gatewayKey", "配置聚合根；schemaVersion 不得高于程序支持版本"),
        ("Resource", "resourceKey", "一对多 Channel；类型为 serial 或 network"),
        ("Channel", "channelKey", "归属 Resource；role 为 collect/forward；协议字段受 role 约束"),
        ("Device", "deviceKey", "归属 Channel；地址、从站号、CA、NodeId 等按协议校验"),
        ("Point", "deviceKey + metric", "全局标识符 metric 唯一策略；协议地址组合需合法"),
        ("ForwardDevice", "deviceKey", "归属转发 Channel；协议决定 UnitID/CA/IED 等字段"),
        ("ForwardPoint", "forwardDevice + metric", "通过 sourceDeviceKey + sourceMetric 引用采集点"),
        ("EdgeComputedPoint", "groupKey + metric", "Inputs 引用采集点或前序派生点，依赖图不得成环"),
        ("User / Role", "username / roleKey", "User 多对一 Role；权限字符串控制 API"),
    ], [1.55, 1.75, 3.20], font_size=8.7)

    heading(doc, "6.2 物理存储设计", 2)
    add_table(doc, ["数据类别", "默认位置/形式", "一致性与容量策略"], [
        ("主配置", "config.local.json", "0600；临时文件 fsync 后 rename；保存前保留 .previous"),
        ("激活凭证", "activation.local.json", "0600；主配置只保留引用，导出时避免泄露密钥"),
        ("AI 凭证", ".runtime/ai-settings.json", "0600；API 只返回掩码；不进入工程导出"),
        ("运行状态", "内存 map + revision", "只保留最新值；重启后由采集重新建立"),
        ("离线缓存", ".runtime/gateway-spool.jsonl 或可移动介质", "按最大字节淘汰最旧记录；恢复后顺序补传"),
        ("历史数据", "可移动介质/.weikong/history/history-YYYYMMDD.jsonl", "按日文件；达到上限删除最旧文件；支持 CSV 导出"),
        ("安全审计", ".runtime/security-audit.jsonl", "0600；最大 2 MB；滚动为 .1；最多读取 500 条"),
        ("临时 PDF", ".runtime/ai-temp", "0600；任务完成、取消、超时或启动时清理"),
    ], [1.25, 2.65, 2.60], font_size=8.8)
    heading(doc, "6.3 JSONL 记录结构", 2)
    add_table(doc, ["记录", "字段"], [
        ("Reading", "time, deviceKey, metrics{metric:value}"),
        ("AuditEvent", "timestamp, username, remoteIp, method, path, status, durationMs, result, detail"),
        ("PointStatus（内存）", "deviceKey, metric, value, updatedAt, error, stale, protocol, address, unit"),
    ], [1.65, 4.85])

    page_break(doc)
    heading(doc, "7. 核心流程与接口设计", 1)
    add_image(doc, diagrams["sequence"])
    add_caption(doc, "图 7-1  采集、计算与上报主时序")
    heading(doc, "7.1 配置保存流程", 2)
    add_number(doc, "客户端读取 /api/config，获得脱敏配置和 ETag。")
    add_number(doc, "提交时携带 If-Match；服务端检测配置是否被其他会话修改。")
    add_number(doc, "解析并恢复掩码密码，执行结构、唯一性、协议和边缘计算校验。")
    add_number(doc, "原子保存配置并保留 previous 备份。")
    add_number(doc, "依次更新边缘程序、runtime、state、转发、缓存、历史和 MQTT；失败则回滚。")
    heading(doc, "7.2 主要接口", 2)
    add_table(doc, ["接口", "方法", "用途", "权限"], [
        ("/api/status", "GET", "聚合运行状态", "status.view"),
        ("/api/ws", "WS", "快照、点位差异、通信状态和命令", "登录会话"),
        ("/api/config", "GET/PUT", "读取/保存完整配置", "config.view/manage"),
        ("/api/config/bootstrap", "GET", "大点表轻量结构", "config.view"),
        ("/api/config/points", "GET", "窗口化点位读取", "config.view"),
        ("/api/config/point-mutations", "POST", "服务端批量复制/修改/删除", "config.manage"),
        ("/api/points/write", "POST", "可写点位下发", "config.manage"),
        ("/api/packets/stream", "GET", "实时报文流", "config.manage"),
        ("/api/ai/chat", "POST/SSE", "网关专用智能对话", "ai.use"),
        ("/api/ai/config-jobs", "POST/GET", "PDF OCR 与配置草稿任务", "ai.use"),
        ("/api/network/*", "GET/POST", "网口、WiFi、蜂窝网络", "network.manage"),
        ("/api/maintenance/*", "POST/GET", "诊断、时间、重启、恢复、审计", "maintenance.run"),
    ], [2.05, 0.75, 2.55, 1.15], font_size=8.2)
    heading(doc, "7.3 WebSocket 消息", 2)
    add_table(doc, ["消息类型", "方向", "说明"], [
        ("snapshot", "服务端→浏览器", "首次连接快照；大点表模式可不携带全量点位"),
        ("gateway.status", "服务端→浏览器", "运行时、资源和汇总状态差异"),
        ("points.diff", "服务端→浏览器", "点位增量；replace 标记可触发替换"),
        ("communications.diff", "服务端→浏览器", "MQTT 等通信连接变化"),
        ("command.execute", "浏览器→服务端", "采集一次、点位下发、连接测试、Ping"),
        ("command.result", "服务端→浏览器", "命令执行结果与 requestId 关联"),
    ], [1.75, 1.35, 3.40])

    page_break(doc)
    heading(doc, "8. 安全、可靠性与性能设计", 1)
    heading(doc, "8.1 安全设计", 2)
    add_table(doc, ["控制点", "设计"], [
        ("身份认证", "本地会话登录；首次账号 admin/123456，要求交付时强制修改"),
        ("授权", "Role + Permission；配置、网络、云、AI、维护等权限分离"),
        ("凭证保护", "配置文件 0600；API 脱敏；AI 和激活凭证独立存放"),
        ("接口保护", "同源 WebSocket、请求体上限、方法限制、安全响应头"),
        ("操作追溯", "非只读请求、登录、AI 生成/应用写入安全审计"),
        ("AI 安全", "只发送脱敏上下文；配置操作先生成确认项；草稿 ETag 防覆盖"),
    ], [1.35, 5.15])
    heading(doc, "8.2 可靠性设计", 2)
    for text in [
        "软件 Supervisor 按健康端点监测进程进度，限制重启频率并支持安全模式文件。",
        "硬件看门狗可独立启停；调试期关闭、交付运行期按设备能力启用。",
        "配置保存采用原子替换和 previous 备份，热加载失败自动恢复。",
        "连接失败采用设备级重连退避；控制写占用连接时采集标记为 deferred，不误报设备离线。",
        "历史与离线缓存仅在介质可用时写入，拔盘不影响主采集。",
    ]: add_bullet(doc, text)
    heading(doc, "8.3 性能与容量", 2)
    add_table(doc, ["场景", "现有策略", "建议验证指标"], [
        ("1,000 点", "完整 WebSocket 快照 + 差异推送", "CPU、RSS、1 秒采集/上报稳定性"),
        ("30,000 点", "虚拟表格、批量读、点位 diff", "滚动帧率、周期抖动、异常率"),
        ("300,000 点配置", "bootstrap/compact、服务端变更、可见窗口值", "首屏时间、保存时间、峰值内存"),
        ("多 MQTT 通道", "每通道独立连接、节拍和状态", "弱 Broker 不影响其他通道"),
    ], [1.35, 3.10, 2.05])

    page_break(doc)
    heading(doc, "9. 原型设计", 1)
    add_body(doc, "原型以“桌面端配置效率 + 移动端查看优先”为原则。桌面端保留左侧资源树、顶部配置/运行切换和单滚动容器；移动端默认折叠配置参数，突出设备状态与实时数据。")
    add_image(doc, diagrams["prototype_status"])
    add_caption(doc, "图 9-1  运行状态页原型")
    add_image(doc, diagrams["prototype_config"])
    add_caption(doc, "图 9-2  智能网关设备与点位配置原型")
    page_break(doc)
    add_image(doc, diagrams["prototype_ai"])
    add_caption(doc, "图 9-3  AI 助手原型")
    add_image(doc, diagrams["prototype_mobile"], width=4.25)
    add_caption(doc, "图 9-4  移动端查看优先原型")
    heading(doc, "9.1 页面信息架构", 2)
    add_table(doc, ["一级模块", "核心任务", "主要页面/弹窗"], [
        ("运行状态", "查看网关、点位、资源和负载", "指标卡、点位表、趋势与错误"),
        ("智能网关", "管理资源、通道、设备和点位", "资源树、设备表单、虚拟点表、协议浏览弹窗"),
        ("边缘计算", "配置输入、表达式和输出", "计算组、输入选择树、CEL 编辑与试算"),
        ("AI 助手", "诊断、对话操作和 PDF 生成", "对话、确认卡、参数表单、草稿差异"),
        ("云平台连接", "激活与 MQTT 通道配置", "JSON 模板编辑、模板预览/实时点位、连接状态"),
        ("网络与维护", "网络、时间、存储、用户和维护", "一致的表单、确认弹窗和加载页"),
    ], [1.35, 2.20, 2.95])
    heading(doc, "9.2 关键交互规范", 2)
    for text in [
        "配置状态不更新实时值并显示“—”；运行状态才启用实时值背景呼吸反馈。",
        "运行状态隐藏复制和删除；协议不支持写入时不显示“下发”。",
        "点表工具条和表头固定，页面只保留一个滚动条；大数据量采用虚拟渲染。",
        "协议浏览统一使用弹窗树，多选后一次添加；OPC UA、IEC 61850、转发点和边缘输入结构一致。",
        "危险操作统一使用全局确认弹窗；重启服务/网关显示不丢失当前页面的加载状态。",
        "手机端把设备信息与配置折叠进摘要卡，搜索和实时数据标题在下滑时吸顶。",
    ]: add_bullet(doc, text)

    page_break(doc)
    heading(doc, "10. 部署与运行设计", 1)
    heading(doc, "10.1 目录建议", 2)
    add_table(doc, ["目录/文件", "用途", "权限建议"], [
        ("/userdata/weikong/bin/weikong-gateway", "主程序", "0755"),
        ("/userdata/weikong/config/config.local.json", "主配置", "0600"),
        ("/userdata/weikong/config/activation.local.json", "激活文件", "0600"),
        ("/userdata/weikong/config/.runtime", "AI、审计、缓存等运行数据", "0700/0600"),
        ("/etc/init.d/S99weikong-gateway", "启动与 supervisor", "0755"),
        ("/media/*/.weikong/history", "可选历史存储", "0700/0600"),
    ], [3.10, 2.35, 1.05])
    heading(doc, "10.2 启动顺序", 2)
    add_number(doc, "加载并校验配置，合并激活文件与默认值。")
    add_number(doc, "Supervisor 根据 WatchdogConfig 决定直接运行或受控运行。")
    add_number(doc, "构建 state、runtime、edgecompute、forward、cloud 和 Web 服务。")
    add_number(doc, "启动健康、心跳、拓扑、发布、边缘计算与采集 worker。")
    add_number(doc, "Web 页面通过 /api/ws 获取初始消息并持续接收差异。")
    heading(doc, "10.3 构建变体", 2)
    add_table(doc, ["变体", "目标", "说明"], [
        ("windows-amd64", "本机功能测试", "不启用 Linux 专属硬件/GOOSE 原始套接字"),
        ("linux-armv7", "EG100 基础版", "适配板载资源、Supervisor、硬件指标"),
        ("linux-armv7 + iec61850_mms", "EG100 IEC61850 版", "CGO 静态链接 libiec61850，提供 MMS 采集/转发"),
    ], [1.70, 1.65, 3.15])

    heading(doc, "11. 验证与验收建议", 1)
    add_table(doc, ["类别", "验收重点"], [
        ("配置", "新增/修改/删除、ETag 冲突、密码掩码保留、热加载失败回滚"),
        ("协议", "各协议连接、批量读取、断线重连、写入、报文监控与转发互通"),
        ("实时", "WebSocket 首次快照、差异推送、重连、慢浏览器不影响采集"),
        ("性能", "1k/30k/300k 点表；CPU、内存、滚动、保存和周期抖动"),
        ("可靠性", "断网、拔盘、Broker 无 ACK、进程假死、硬件看门狗和安全模式"),
        ("安全", "登录限速、权限矩阵、密钥不出日志、审计、工程导出脱敏"),
        ("AI", "错误凭证、PDF 超限、未知字段不猜测、草稿冲突、确认与回滚"),
    ], [1.20, 5.30])

    heading(doc, "附录 A. 术语与代码映射", 1)
    add_table(doc, ["术语", "含义", "代码位置"], [
        ("采集 lane", "同设备/通道的一组周期任务", "internal/runtime/runtime.go"),
        ("最新值状态", "当前值、错误、更新时间和 stale", "internal/state/state.go"),
        ("离线 Spool", "MQTT 失败时的顺序 JSONL 缓存", "internal/cache/spool.go"),
        ("历史 Store", "按日 JSONL 历史和 CSV 导出", "internal/history/store.go"),
        ("数据管理器", "WebSocket 差异与日志消息", "internal/datamanager/manager.go"),
        ("配置聚合根", "资源、通道、设备、点位及安全等", "internal/config/config.go"),
        ("前端壳", "导航、资源树、运行/配置切换", "webui/src/App.vue"),
    ], [1.35, 2.70, 2.45], font_size=8.8)
    add_callout(doc, "实现基线提示", "本文档反映 2026-08-20 工作区实现。后续新增协议、数据库或云端控制面时，应同步更新架构图、类图、ER 图、接口表和原型。")

    # Document properties.
    props = doc.core_properties
    props.title = "微控 Go 工业网关系统设计说明书"
    props.subject = "系统总体结构、类图、数据库设计与原型设计"
    props.author = "微控科技"
    props.keywords = "Go, EG100, 工业网关, 系统设计, IEC61850, MQTT"
    props.comments = "Generated from the current gateway implementation baseline."
    doc.save(OUT_FILE)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    diagrams = {
        "architecture": save_architecture(),
        "deployment": save_deployment(),
        "class_core": save_class_core(),
        "class_domain": save_class_domain(),
        "er": save_er(),
        "sequence": save_sequence(),
        "prototype_status": save_prototype_status(),
        "prototype_config": save_prototype_config(),
        "prototype_ai": save_prototype_ai(),
        "prototype_mobile": save_prototype_mobile(),
    }
    build_document(diagrams)
    print(OUT_FILE)


if __name__ == "__main__":
    main()
